"""Convert Online Retail II XLSX sheets into a unified, high-performance CSV.

Outputs:
- data/raw/online_retail_raw.csv (~1,067,371 rows)
"""
import csv
import sys
from pathlib import Path
from openpyxl import load_workbook

raw_dir = Path(__file__).resolve().parents[1] / "data" / "raw"
xlsx_file = raw_dir / "online_retail_II.xlsx"
csv_file = raw_dir / "online_retail_raw.csv"

if csv_file.exists() and csv_file.stat().st_size > 10_000_000:
    print(f"Clean CSV already exists at {csv_file}")
    sys.exit(0)

print(f"Reading {xlsx_file} in streaming read-only mode...")
wb = load_workbook(str(xlsx_file), read_only=True, data_only=True)

total_written = 0
with open(csv_file, "w", newline="", encoding="utf-8") as f_out:
    writer = csv.writer(f_out)
    header_written = False

    for sheet_name in ["Year 2009-2010", "Year 2010-2011"]:
        print(f"Processing sheet {sheet_name}...")
        ws = wb[sheet_name]
        is_first = True
        for row in ws.iter_rows(values_only=True):
            if is_first:
                is_first = False
                if not header_written:
                    writer.writerow(["Invoice", "StockCode", "Description", "Quantity", "InvoiceDate", "Price", "Customer ID", "Country"])
                    header_written = True
                continue
            if row[0] is None:
                continue
            writer.writerow(row)
            total_written += 1
            if total_written % 200_000 == 0:
                print(f"  ... written {total_written:,} rows")

print(f"Finished! Written {total_written:,} transactions to {csv_file}")
